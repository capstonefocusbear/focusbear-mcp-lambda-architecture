#!/usr/bin/env python3
"""Create a concise client-facing PowerPoint for the Focus Bear MCP proposal."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "pptx"
PPTX_OUT = OUT / "Focus_Bear_AWS_Lambda_MCP_Architecture_Proposal.pptx"

NAVY = "0B2545"
BLUE = "2E74B5"
TEAL = "1B998B"
SKY = "E8F1FA"
MINT = "E5F5F1"
AMBER = "FDF0D5"
ROSE = "FBE7E8"
SLATE = "52657A"
LIGHT = "F6F8FB"
WHITE = "FFFFFF"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = NAVY,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    font: str = "Aptos",
    margin: float = 0.04,
    valign=MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return shape


def add_box(slide, x, y, w, h, *, fill=WHITE, line=SKY, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(1)
    return shape


def add_card(slide, title, body, x, y, w, h, *, fill=WHITE, accent=BLUE, body_size=14):
    add_box(slide, x, y, w, h, fill=fill, line=accent)
    accent_bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = rgb(accent)
    accent_bar.line.fill.background()
    add_text(slide, title, x + 0.24, y + 0.16, w - 0.36, 0.35, size=17, color=NAVY, bold=True)
    add_text(slide, body, x + 0.24, y + 0.62, w - 0.38, h - 0.72, size=body_size, color=SLATE)


def add_bullets(slide, items, x, y, w, h, *, size=16, color=SLATE, bullet_color=TEAL, gap=0.44):
    for index, item in enumerate(items):
        current_y = y + index * gap
        add_text(slide, "•", x, current_y, 0.22, 0.25, size=size + 2, color=bullet_color, bold=True)
        add_text(slide, item, x + 0.23, current_y, w - 0.23, 0.35, size=size, color=color)


def add_arrow(slide, x1, y1, x2, y2, *, color=BLUE, width=1.6):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    line.line.end_arrowhead = True


def add_header(slide, eyebrow: str, title: str, subtitle: str = ""):
    add_text(slide, eyebrow.upper(), 0.62, 0.34, 4.8, 0.28, size=10.5, color=TEAL, bold=True)
    add_text(slide, title, 0.58, 0.66, 12.05, 0.62, size=27, color=NAVY, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.61, 1.30, 11.85, 0.34, size=13, color=SLATE)
    rule = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.60), Inches(1.73), Inches(12.15), Inches(0.025))
    rule.fill.solid()
    rule.fill.fore_color.rgb = rgb(SKY)
    rule.line.fill.background()


def add_footer(slide, page: int):
    rule = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.60), Inches(7.12), Inches(12.15), Inches(0.015))
    rule.fill.solid()
    rule.fill.fore_color.rgb = rgb(SKY)
    rule.line.fill.background()
    add_text(slide, "FOCUS BEAR  |  AWS LAMBDA MCP PROPOSAL  |  CLIENT DRAFT", 0.60, 7.19, 6.5, 0.20, size=8.5, color=SLATE, bold=True)
    add_text(slide, str(page), 12.18, 7.17, 0.55, 0.20, size=9, color=SLATE, align=PP_ALIGN.RIGHT)


def process_step(slide, title, detail, x, y, fill):
    add_box(slide, x, y, 2.15, 1.07, fill=fill, line=BLUE)
    add_text(slide, title, x + 0.12, y + 0.15, 1.91, 0.25, size=15.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, detail, x + 0.13, y + 0.50, 1.89, 0.35, size=11.5, color=SLATE, align=PP_ALIGN.CENTER)


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]

    # 1. Cover
    slide = deck.slides.add_slide(blank)
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(NAVY)
    slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ARC, Inches(9.5), Inches(-1.55), Inches(5.2), Inches(5.2)).line.color.rgb = rgb(TEAL)
    slide.shapes[-1].fill.background()
    add_text(slide, "FOCUS BEAR", 0.80, 0.83, 3.0, 0.30, size=13, color="70D5C8", bold=True)
    add_text(slide, "AWS Lambda MCP\nArchitecture Proposal", 0.76, 1.38, 8.0, 1.58, size=33, color=WHITE, bold=True)
    add_text(slide, "A secure, scalable way to give AI agents scoped access to Focus Bear tasks.", 0.80, 3.23, 7.7, 0.48, size=17, color="D8E7F5")
    add_box(slide, 0.80, 4.25, 3.25, 0.68, fill="173C63", line="2A5E8E")
    add_text(slide, "Client presentation  |  11 Sep 2026", 0.98, 4.47, 2.9, 0.24, size=11.5, color=WHITE, bold=True)
    add_text(slide, "Draft for discussion", 0.81, 6.85, 2.2, 0.22, size=10, color="9AB8D1")

    # 2. Recommendation
    slide = deck.slides.add_slide(blank)
    add_header(slide, "Executive decision", "Recommended approach", "Start with a thin, stateless MCP adapter; keep Focus Bear’s API and database as the source of truth.")
    add_card(slide, "Public edge", "API Gateway REST API\n• HTTPS endpoint\n• rate limits and optional WAF\n• response streaming when needed", 0.66, 2.02, 3.85, 2.15, fill=SKY, accent=BLUE, body_size=14)
    add_card(slide, "MCP adapter", "AWS Lambda\n• validates MCP protocol\n• maps tools to API requests\n• does not own business data", 4.74, 2.02, 3.85, 2.15, fill=MINT, accent=TEAL, body_size=14)
    add_card(slide, "Focus Bear authority", "Existing NestJS API + PostgreSQL\n• validates user token and scopes\n• preserves agent/task assignment\n• applies domain rules", 8.82, 2.02, 3.85, 2.15, fill="EDF2F7", accent=NAVY, body_size=14)
    add_box(slide, 0.67, 4.60, 12.0, 1.43, fill=LIGHT, line=SKY)
    add_text(slide, "Why this is the right v1", 0.93, 4.83, 2.5, 0.28, size=16.5, color=NAVY, bold=True)
    add_bullets(slide, ["Fast to ship: reuses the existing external-MCP endpoints and permission checks.", "Low operational overhead: no new database, queue, or always-on service.", "Safe evolution path: add OAuth, more tools, or container hosting only when usage requires it."], 3.26, 4.74, 8.95, 0.88, size=13.4, gap=0.28)
    add_footer(slide, 2)

    # 3. Current fit
    slide = deck.slides.add_slide(blank)
    add_header(slide, "Current-state fit", "It builds on the architecture already in Focus Bear", "The repository already has an external MCP API boundary and a Lambda-oriented MCP server prototype.")
    process_step(slide, "MCP client", "Presents a token\nand asks for a tool", 0.70, 2.40, SKY)
    process_step(slide, "AWS Lambda", "Protocol adapter\nand request mapper", 3.18, 2.40, MINT)
    process_step(slide, "Focus Bear API", "NestJS controllers,\nguards and services", 5.66, 2.40, "EAF0F6")
    process_step(slide, "Task rules", "Scope + agent\nassignment checks", 8.14, 2.40, AMBER)
    process_step(slide, "PostgreSQL", "Tasks, projects\nand token records", 10.62, 2.40, "F3F5F7")
    for x in (2.85, 5.33, 7.81, 10.29):
        add_arrow(slide, x, 2.93, x + 0.30, 2.93)
    add_card(slide, "Existing building blocks", "• External MCP task endpoints\n• Bearer-token validation and scopes\n• Token-to-agent assignment on tasks\n• Per-token throttling\n• A Lambda streaming prototype", 0.80, 4.48, 5.85, 1.95, fill=SKY, accent=BLUE, body_size=13)
    add_card(slide, "Important v1 cleanup", "• Remove development fallback credentials\n• Restore real authentication on token issuance\n• Make the internal service key a required secret\n• Agree an explicit public endpoint / network path", 6.90, 4.48, 5.55, 1.95, fill=ROSE, accent="C9565A", body_size=13)
    add_footer(slide, 3)

    # 4. Architecture overview
    slide = deck.slides.add_slide(blank)
    add_header(slide, "Target architecture", "One clear request path, with two trust boundaries", "The Lambda layer is deliberately thin; the Focus Bear API remains the policy enforcement point.")
    process_step(slide, "AI agent", "MCP client\nBearer token", 0.65, 2.45, SKY)
    process_step(slide, "API Gateway", "REST API\n/ mcp", 3.12, 2.45, "DCEBFA")
    process_step(slide, "Lambda adapter", "MCP protocol\ninput validation", 5.59, 2.45, MINT)
    process_step(slide, "Focus Bear API", "Internal key +\nuser token", 8.06, 2.45, "EAF0F6")
    process_step(slide, "Data", "PostgreSQL\nexisting schema", 10.53, 2.45, AMBER)
    for x in (2.80, 5.27, 7.74, 10.21):
        add_arrow(slide, x, 2.98, x + 0.32, 2.98)
    add_box(slide, 5.48, 4.35, 2.40, 0.95, fill=LIGHT, line=TEAL)
    add_text(slide, "Secrets Manager", 5.70, 4.56, 1.98, 0.20, size=14.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "rotated internal credential", 5.62, 4.84, 2.14, 0.18, size=10.5, color=SLATE, align=PP_ALIGN.CENTER)
    add_arrow(slide, 6.68, 4.35, 6.68, 3.62, color=TEAL)
    add_box(slide, 8.38, 4.35, 2.42, 0.95, fill=LIGHT, line=BLUE)
    add_text(slide, "CloudWatch + X-Ray", 8.50, 4.56, 2.18, 0.20, size=14, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "safe logs, traces, alarms", 8.52, 4.84, 2.14, 0.18, size=10.5, color=SLATE, align=PP_ALIGN.CENTER)
    add_arrow(slide, 9.58, 4.35, 9.58, 3.62, color=BLUE)
    add_text(slide, "Trust boundary 1: public MCP client → AWS\nTrust boundary 2: AWS adapter → Focus Bear API", 0.92, 5.75, 5.8, 0.55, size=14, color=SLATE)
    add_footer(slide, 4)

    # 5. Options
    slide = deck.slides.add_slide(blank)
    add_header(slide, "Options and trade-offs", "Hosting choices", "Recommendation is driven by MCP transport compatibility and low operational burden, not technology novelty.")
    add_card(slide, "Recommended", "API Gateway REST API + Lambda\n\nPros: supports Lambda response streaming; managed edge controls; serverless operations.\n\nTrade-off: slightly more configuration than HTTP API.", 0.65, 2.04, 3.00, 3.85, fill=MINT, accent=TEAL, body_size=13.2)
    add_card(slide, "Conditional fallback", "API Gateway HTTP API + Lambda\n\nPros: simpler and often lower cost.\n\nTrade-off: use only if all target MCP clients work with buffered HTTP responses; not the streaming default.", 3.88, 2.04, 3.00, 3.85, fill=SKY, accent=BLUE, body_size=13.2)
    add_card(slide, "Not for v1", "Lambda Function URL\n\nPros: least infrastructure.\n\nTrade-off: fewer gateway controls and less deliberate public edge; not preferred for a client-facing service.", 7.11, 2.04, 2.68, 3.85, fill=AMBER, accent="C99127", body_size=13.0)
    add_card(slide, "Later, if needed", "ECS / Fargate\n\nPros: best for persistent connections, very high sustained usage, or complex connection state.\n\nTrade-off: more operations and cost; unnecessary for a focused v1.", 10.02, 2.04, 2.68, 3.85, fill=LIGHT, accent=NAVY, body_size=13.0)
    add_footer(slide, 5)

    # 6. Security
    slide = deck.slides.add_slide(blank)
    add_header(slide, "Security model", "Two credentials, two different purposes", "A user token answers “what may this agent do?”; an internal key answers “is this trusted infrastructure?”")
    add_card(slide, "1. User MCP token", "Passed from MCP client to Lambda and forwarded to the Focus Bear API.\n\nValidated as a raw opaque token; scope and agent assignment are checked server-side.", 0.72, 2.10, 3.70, 2.34, fill=SKY, accent=BLUE, body_size=14)
    add_card(slide, "2. Internal service key", "Stored in AWS Secrets Manager and read by the Lambda execution role.\n\nProves the request came through the approved adapter; rotate it without redeploying clients.", 4.80, 2.10, 3.70, 2.34, fill=MINT, accent=TEAL, body_size=14)
    add_card(slide, "3. Existing policy checks", "Focus Bear API enforces token scope, user ownership, and the token/agent assignment already recorded on a task.\n\nNo direct database access from Lambda.", 8.88, 2.10, 3.70, 2.34, fill="EAF0F6", accent=NAVY, body_size=14)
    add_box(slide, 0.72, 4.87, 11.86, 0.98, fill=ROSE, line="C9565A")
    add_text(slide, "Security gate before production", 0.98, 5.10, 3.0, 0.25, size=15.5, color="8D3033", bold=True)
    add_text(slide, "Eliminate hard-coded development fallbacks, require real authentication for token issuance, redact secrets in logs, and set explicit rate / concurrency limits.", 3.60, 5.08, 8.48, 0.38, size=13.3, color="8D3033")
    add_footer(slide, 6)

    # 7. V1 capabilities
    slide = deck.slides.add_slide(blank)
    add_header(slide, "V1 capability", "A focused tool set keeps risk and delivery time low", "Start with existing task operations; every write goes through the same Focus Bear API validation path.")
    tools = [
        ("List tasks", "Query tasks assigned to the authenticated agent", SKY, BLUE),
        ("Get task", "Read a permitted task by ID", MINT, TEAL),
        ("Update status", "Move a permitted task to an allowed status", "EAF0F6", NAVY),
        ("Add note", "Append a task note through the API", AMBER, "C99127"),
        ("Project statuses", "List valid statuses for a project", LIGHT, SLATE),
    ]
    for index, (name, detail, fill, accent) in enumerate(tools):
        x = 0.75 + index * 2.45
        add_card(slide, name, detail, x, 2.36, 2.16, 2.18, fill=fill, accent=accent, body_size=12.6)
    add_box(slide, 0.75, 5.18, 11.90, 0.84, fill=LIGHT, line=SKY)
    add_text(slide, "Principle: expose only narrow, named business actions—not generic database access or arbitrary API proxying.", 0.99, 5.45, 11.33, 0.25, size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 7)

    # 8. Operations
    slide = deck.slides.add_slide(blank)
    add_header(slide, "Reliability and operations", "Designed for predictable behaviour under load", "The adapter should fail safely and make both client and engineering diagnosis straightforward.")
    add_card(slide, "Capacity controls", "Reserved concurrency protects the Focus Bear API; API Gateway throttling limits abusive traffic; short timeouts prevent request pile-up.", 0.70, 2.10, 3.70, 2.35, fill=SKY, accent=BLUE, body_size=14)
    add_card(slide, "Error behaviour", "Return protocol-safe errors. Do not retry writes automatically unless an idempotency key is in place. Preserve a correlation ID for support.", 4.80, 2.10, 3.70, 2.35, fill=MINT, accent=TEAL, body_size=14)
    add_card(slide, "Observability", "Track invocation volume, latency, error rate, throttles and cold starts. Alert on sustained failures; never log raw credentials or full task content by default.", 8.90, 2.10, 3.70, 2.35, fill="EAF0F6", accent=NAVY, body_size=14)
    add_box(slide, 0.70, 4.92, 11.90, 0.90, fill=AMBER, line="C99127")
    add_text(slide, "Cost profile", 0.96, 5.16, 1.55, 0.25, size=15.5, color="6D4F08", bold=True)
    add_text(slide, "Pay per request and execution time; no idle service cost. The main cost variables are request volume, streaming duration, logs/traces, and Secrets Manager use.", 2.40, 5.15, 9.72, 0.30, size=13.4, color="6D4F08")
    add_footer(slide, 8)

    # 9. Delivery
    slide = deck.slides.add_slide(blank)
    add_header(slide, "Delivery path", "A staged rollout makes the risk visible and manageable", "Exact duration follows the answers to the decision questions on the next slide.")
    phases = [
        ("1", "Confirm contract", "Client support, tool schemas, SLOs"),
        ("2", "Harden API", "Remove fallbacks, auth token issuance"),
        ("3", "Build adapter", "MCP handlers, validation, errors"),
        ("4", "Provision AWS", "Gateway, Lambda, secrets, IAM"),
        ("5", "Verify", "Unit, integration, security, load"),
        ("6", "Pilot", "Limited agents, metrics, feedback"),
    ]
    for index, (number, title, detail) in enumerate(phases):
        x = 0.67 + index * 2.04
        circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.72), Inches(2.18), Inches(0.58), Inches(0.58))
        circle.fill.solid(); circle.fill.fore_color.rgb = rgb(TEAL if index < 2 else BLUE)
        circle.line.fill.background()
        add_text(slide, number, x + 0.72, 2.30, 0.58, 0.20, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, title, x, 2.98, 2.0, 0.30, size=14.2, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.05, 3.36, 1.90, 0.55, size=11.3, color=SLATE, align=PP_ALIGN.CENTER)
        if index < len(phases) - 1:
            add_arrow(slide, x + 1.54, 2.47, x + 2.08, 2.47, color="91ABC3")
    add_box(slide, 0.83, 4.70, 11.65, 1.10, fill=LIGHT, line=SKY)
    add_text(slide, "Launch guardrail", 1.10, 5.00, 1.72, 0.25, size=15.5, color=NAVY, bold=True)
    add_text(slide, "Production launch requires passed security checks, successful end-to-end MCP client tests, agreed error and support handling, and monitoring / rollback readiness.", 2.82, 4.98, 8.96, 0.38, size=13.4, color=SLATE)
    add_footer(slide, 9)

    # 10. Decisions
    slide = deck.slides.add_slide(blank)
    add_header(slide, "Decision points", "What we need to confirm with Focus Bear", "These answers turn the proposal into an implementation plan and a reliable estimate.")
    questions = [
        ("Target clients", "Which MCP clients must be supported—and do any need streamed responses or server-initiated messages?"),
        ("Network path", "Where is the production Focus Bear API reachable from AWS, and how will the internal service credential be rotated?"),
        ("Usage profile", "Expected agents, peak concurrent calls, request/response sizes, and target latency / availability?"),
        ("Identity model", "Are opaque MCP tokens sufficient for v1, or is third-party OAuth / delegated identity required now?"),
    ]
    for index, (title, detail) in enumerate(questions):
        y = 2.05 + index * 1.03
        add_box(slide, 0.75, y, 11.85, 0.78, fill=SKY if index % 2 == 0 else LIGHT, line=SKY, radius=True)
        add_text(slide, title, 1.02, y + 0.18, 1.80, 0.24, size=14.5, color=NAVY, bold=True)
        add_text(slide, detail, 2.84, y + 0.16, 9.30, 0.34, size=13.2, color=SLATE)
    add_text(slide, "Proposed next step: validate client transport requirements, then run a short technical discovery to lock the interfaces and launch controls.", 0.83, 6.37, 11.62, 0.35, size=15.5, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, 10)

    deck.save(PPTX_OUT)
    print(PPTX_OUT)


if __name__ == "__main__":
    build()
