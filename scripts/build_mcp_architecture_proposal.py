#!/usr/bin/env python3
"""Build the client-facing Focus Bear AWS Lambda MCP architecture proposal.

The source of truth is docs/architecture/Focusware_AWS_Lambda_MCP_Architecture_Proposal.md.
This generator creates a styled DOCX and a diagram asset for document rendering.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/architecture/Focusware_AWS_Lambda_MCP_Architecture_Proposal.md"
OUT = ROOT / "output"
DOCX_OUT = OUT / "docx" / "Focus_Bear_AWS_Lambda_MCP_Architecture_Proposal.docx"
ASSETS = OUT / "assets"
DIAGRAM = ASSETS / "focus_bear-mcp-architecture.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5E6C84"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for old in list(grid):
        grid.remove(old)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    node = OxmlElement("w:keepNext")
    p_pr.append(node)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_font(run, size=11, color=INK, bold=None, italic=None, name="Calibri") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, text: str, size=11, color=INK) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, color=color, bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size, color=color)


def paragraph_format(paragraph, before=0, after=6, line=1.1, align=None) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_body(doc, text: str) -> None:
    p = doc.add_paragraph()
    paragraph_format(p, after=6, line=1.1)
    add_inline(p, text)


def add_heading(doc, text: str, level: int) -> None:
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    p = doc.add_paragraph(style=style)
    paragraph_format(p, before={1: 16, 2: 12, 3: 8}[level], after={1: 8, 2: 6, 3: 4}[level], line=1.0)
    add_inline(p, text, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level])
    keep_with_next(p)


def add_bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    paragraph_format(p, after=4, line=1.167)
    add_inline(p, text)


def add_numbered(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    paragraph_format(p, after=4, line=1.167)
    add_inline(p, text)


def restart_numbering(doc, paragraph) -> None:
    """Give a numbered list that begins at 1 its own real Word numbering instance."""
    numbering = doc.part.numbering_part.element
    style_num_pr = doc.styles["List Number"]._element.pPr.numPr
    base_num_id = style_num_pr.numId.val
    existing = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    new_num_id = max(existing) + 1
    base = next(node for node in numbering.findall(qn("w:num")) if int(node.get(qn("w:numId"))) == int(base_num_id))
    abstract_id = base.find(qn("w:abstractNumId")).get(qn("w:val"))
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    ilvl = num_pr.get_or_add_ilvl()
    ilvl.set(qn("w:val"), "0")
    num_id = num_pr.get_or_add_numId()
    num_id.set(qn("w:val"), str(new_num_id))


def add_code(doc, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        paragraph_format(p, before=0, after=0, line=1.0)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.right_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), LIGHT_GRAY)
        p_pr.append(shd)
        run = p.add_run(line)
        set_font(run, size=8.5, color=INK, name="Courier New")


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    if cols == 2:
        widths = [2700, 6660]
    elif cols == 3:
        widths = [2100, 3630, 3630]
    elif cols == 4:
        widths = [1750, 2550, 2550, 2510]
    else:
        widths = [9360 // cols for _ in range(cols)]
        widths[-1] += 9360 - sum(widths)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for r_i, row_data in enumerate(rows):
        row = table.rows[r_i]
        prevent_row_split(row)
        if r_i == 0:
            set_repeat_table_header(row)
        for c_i, text in enumerate(row_data):
            cell = row.cells[c_i]
            if r_i == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            paragraph_format(p, after=0, line=1.0)
            add_inline(p, text, size=9.2, color=INK)
            if r_i == 0:
                for run in p.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    paragraph_format(spacer, after=3, line=1.0)


def font(size: int, bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def draw_box(draw, bounds, title, lines, fill, outline=INK):
    def colour(value: str) -> str:
        return f"#{value}" if len(value) == 6 else value

    x1, y1, x2, y2 = bounds
    draw.rounded_rectangle(bounds, radius=18, fill=colour(fill), outline=colour(outline), width=4)
    title_font = font(28, bold=True)
    body_font = font(20)
    title_bbox = draw.textbbox((0, 0), title, font=title_font)
    draw.text((x1 + (x2 - x1 - (title_bbox[2] - title_bbox[0])) / 2, y1 + 22), title, font=title_font, fill=colour(INK))
    y = y1 + 70
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=body_font)
        draw.text((x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2, y), line, font=body_font, fill=colour(INK))
        y += 28


def arrow(draw, start, end, label=None):
    draw.line([start, end], fill=f"#{DARK_BLUE}", width=5)
    x, y = end
    dx, dy = end[0] - start[0], end[1] - start[1]
    if abs(dx) >= abs(dy):
        points = [(x, y), (x - 18 if dx > 0 else x + 18, y - 10), (x - 18 if dx > 0 else x + 18, y + 10)]
    else:
        points = [(x, y), (x - 10, y - 18 if dy > 0 else y + 18), (x + 10, y - 18 if dy > 0 else y + 18)]
    draw.polygon(points, fill=f"#{DARK_BLUE}")
    if label:
        f = font(17)
        bx = (start[0] + end[0]) / 2
        by = (start[1] + end[1]) / 2 - 30
        bb = draw.textbbox((0, 0), label, font=f)
        draw.rounded_rectangle((bx - (bb[2]-bb[0])/2 - 5, by - 3, bx + (bb[2]-bb[0])/2 + 5, by + 24), radius=4, fill=f"#{WHITE}")
        draw.text((bx - (bb[2]-bb[0])/2, by), label, font=f, fill=f"#{DARK_BLUE}")


def create_diagram() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (2400, 1180), "white")
    draw = ImageDraw.Draw(image)
    def rgb(value: str) -> str:
        return f"#{value}" if len(value) == 6 else value

    draw.text((70, 35), "Focus Bear AWS Lambda MCP request path and trust boundaries", font=font(38, True), fill=rgb(INK))
    draw.rounded_rectangle((45, 120, 2355, 810), radius=24, outline="#B7C3D6", width=4)
    draw.text((80, 140), "Public client and AWS MCP edge", font=font(22, True), fill=rgb(MUTED))
    draw_box(draw, (100, 330, 420, 530), "MCP client", ["AI agent / desktop", "Bearer MCP token"], "E8EEF5")
    draw_box(draw, (550, 300, 910, 560), "API Gateway", ["REST API", "GET + POST /mcp", "throttle + WAF"], "DDEBF7")
    draw_box(draw, (1050, 300, 1410, 560), "Lambda adapter", ["Streamable HTTP", "tool validation", "error mapping"], "CFE2F3")
    draw_box(draw, (1560, 300, 1950, 560), "Focus Bear API", ["NestJS ExternalMcp", "token + scope checks", "domain services"], "E2F0D9")
    draw_box(draw, (2050, 330, 2320, 530), "PostgreSQL", ["tokens, tasks", "projects"], "FFF2CC")
    arrow(draw, (420, 430), (550, 430), "HTTPS JSON-RPC")
    arrow(draw, (910, 430), (1050, 430), "proxy invocation")
    arrow(draw, (1410, 430), (1560, 430), "HTTPS + bearer + internal key")
    arrow(draw, (1950, 430), (2050, 430), "TypeORM")
    draw.rounded_rectangle((1010, 870, 1450, 1070), radius=18, fill="#F4F6F9", outline=rgb(INK), width=4)
    draw.text((1100, 898), "Secrets Manager", font=font(26, True), fill=rgb(INK))
    draw.text((1084, 940), "rotated internal credential", font=font(19), fill=rgb(INK))
    draw.text((1114, 970), "least-privilege IAM", font=font(19), fill=rgb(INK))
    arrow(draw, (1230, 870), (1230, 560), "cached secret")
    draw.rounded_rectangle((1570, 870, 2030, 1070), radius=18, fill="#F4F6F9", outline=rgb(INK), width=4)
    draw.text((1640, 898), "CloudWatch / X-Ray", font=font(26, True), fill=rgb(INK))
    draw.text((1645, 940), "safe logs, metrics", font=font(19), fill=rgb(INK))
    draw.text((1660, 970), "traces and alarms", font=font(19), fill=rgb(INK))
    arrow(draw, (1800, 870), (1800, 560), "telemetry")
    image.save(DIAGRAM)


def parse_markdown(doc: Document, content: str) -> None:
    lines = content.splitlines()
    i = 0
    first_h1_skipped = False
    while i < len(lines):
        line = lines[i]
        if line.startswith("```mermaid"):
            while i < len(lines) and not (i > 0 and lines[i] == "```"):
                i += 1
            doc.add_picture(str(DIAGRAM), width=Inches(6.5))
            i += 1
            continue
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and lines[i] != "```":
                code_lines.append(lines[i])
                i += 1
            add_code(doc, code_lines)
            i += 1
            continue
        if line.startswith("# "):
            if not first_h1_skipped:
                first_h1_skipped = True
            else:
                add_heading(doc, line[2:], 1)
            i += 1
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:], 1)
            i += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:], 2)
            i += 1
            continue
        if line.startswith("#### "):
            add_heading(doc, line[5:], 3)
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = []
            for table_line in table_lines:
                cells = [x.strip() for x in table_line.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            continue
        if line.startswith("- "):
            add_bullet(doc, line[2:])
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            paragraph_format(p, after=4, line=1.167)
            add_inline(p, re.sub(r"^\d+\. ", "", line))
            i += 1
            continue
        if line.strip() == "":
            i += 1
            continue
        # Markdown hard line-breaks are presentation-only; retain the content.
        add_body(doc, line.rstrip("  "))
        i += 1


def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK_BLUE)]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def add_page_number(paragraph) -> None:
    run = paragraph.add_run("Page ")
    set_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_cover(doc: Document) -> None:
    for _ in range(7):
        spacer = doc.add_paragraph()
        paragraph_format(spacer, after=0, line=1.0)
    kicker = doc.add_paragraph()
    paragraph_format(kicker, after=14, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = kicker.add_run("FOCUS BEAR")
    set_font(run, size=12, color=BLUE, bold=True)
    title = doc.add_paragraph()
    paragraph_format(title, after=8, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = title.add_run("AWS Lambda MCP\nArchitecture Proposal")
    set_font(run, size=28, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    paragraph_format(subtitle, after=28, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = subtitle.add_run("Implementation-ready design for a secure, scoped task MCP server")
    set_font(run, size=14, color=MUTED)
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2500, 6860])
    data = [
        ("Project", "Focus Bear MCP server"),
        ("Status", "Draft for client and engineering review"),
        ("Prepared by", "Focus Bear Engineering"),
        ("Date", "11 September 2026"),
    ]
    for idx, (label, value) in enumerate(data):
        set_cell_shading(table.rows[idx].cells[0], DARK_BLUE)
        p = table.rows[idx].cells[0].paragraphs[0]
        paragraph_format(p, after=0, line=1.0)
        r = p.add_run(label)
        set_font(r, size=10.5, color=WHITE, bold=True)
        p = table.rows[idx].cells[1].paragraphs[0]
        paragraph_format(p, after=0, line=1.0)
        r = p.add_run(value)
        set_font(r, size=10.5, color=INK)
    doc.add_page_break()


def build() -> None:
    OUT.mkdir(exist_ok=True)
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    create_diagram()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    setup_styles(doc)

    header = section.header.paragraphs[0]
    paragraph_format(header, after=0, line=1.0)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("FOCUS BEAR AWS LAMBDA MCP ARCHITECTURE PROPOSAL")
    set_font(run, size=8.5, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    paragraph_format(footer, after=0, line=1.0)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Draft for client and engineering review | ")
    set_font(run, size=8.5, color=MUTED)
    add_page_number(footer)

    add_cover(doc)
    content = SOURCE.read_text()
    # The source draft predates the confirmed company name. Keep its technical
    # substance while ensuring all client-facing output uses the correct brand.
    content = content.replace("Focusware (Focus Bear backend)", "Focus Bear backend")
    content = (
        content.replace("FOCUSWARE_", "FOCUS_BEAR_")
        .replace("Focusware", "Focus Bear")
        .replace("focusware", "focusbear")
    )
    parse_markdown(doc, content)
    doc.save(DOCX_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    build()
